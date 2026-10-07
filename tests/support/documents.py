"""Generated document fixtures for the rendering tests — bytes built in process, no files.

Every builder returns the bytes of a real file of its format, made with
the same library the renderer reads it with, so a test states what the
document holds and asserts what the rendering says about it.
"""

from __future__ import annotations

import struct
import zlib
from io import BytesIO
from typing import TYPE_CHECKING

import docx
import openpyxl
import pptx
import pypdf
from pptx.util import Inches
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

if TYPE_CHECKING:
    from collections.abc import Mapping, Sequence


def png_bytes(width: int = 2, height: int = 2) -> bytes:
    """A valid solid-red PNG of the given size."""
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def pdf_bytes(pages: Sequence[str | None]) -> bytes:
    """One page per item: its lines of Helvetica text, or a blank page for ``None``."""
    writer = pypdf.PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        if text is None:
            continue
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        fonts = DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
        shown = " ".join(f"({line}) Tj T*" for line in text.split("\n"))
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td 14 TL {shown} ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def surrogate_pdf_bytes() -> bytes:
    """A one-page PDF whose ToUnicode map sends the glyph ``A`` to a lone surrogate (U+D800)."""
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    cmap = DecodedStreamObject()
    cmap.set_data(
        b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap "
        b"/CMapName /Custom def 1 begincodespacerange <00> <FF> endcodespacerange "
        b"1 beginbfchar <41> <D800> endbfchar endcmap CMapName currentdict /CMap defineresource pop end end"
    )
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/ToUnicode"): writer._add_object(cmap),
        }
    )
    fonts = DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (A) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def encrypted_pdf_bytes(pages: Sequence[str | None], password: str = "secret") -> bytes:
    """*pages* as a PDF that opens only with *password* (an empty one when ``password`` is empty)."""
    writer = pypdf.PdfWriter(clone_from=pypdf.PdfReader(BytesIO(pdf_bytes(pages))))
    writer.encrypt(password, owner_password="owner")
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------


def docx_bytes(
    paragraphs: Sequence[str] = (),
    *,
    title: str | None = None,
    headings: Sequence[tuple[int, str]] = (),
    table: Sequence[Sequence[str]] | None = None,
    header: str | None = None,
    footer: str | None = None,
    comment: str | None = None,
    picture: bool = False,
    second_section: bool = False,
) -> bytes:
    """A document in this body order: title, headings, paragraphs, table, then a picture.

    *comment* is attached to the first paragraph's runs; *header* and
    *footer* fill the first section's header and footer; with
    *second_section* a second section follows, its header and footer
    linked to the first's.
    """
    document = docx.Document()
    if title is not None:
        document.add_heading(title, 0)
    for level, text in headings:
        document.add_heading(text, level)
    written = [document.add_paragraph(text) for text in paragraphs]
    if comment is not None and written:
        document.add_comment(written[0].runs, text=comment, author="Reviewer")
    if table is not None:
        grid = document.add_table(rows=len(table), cols=max(len(row) for row in table))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    if picture:
        document.add_picture(BytesIO(png_bytes()))
    section = document.sections[0]
    if header is not None:
        section.header.paragraphs[0].text = header
    if footer is not None:
        section.footer.paragraphs[0].text = footer
    if second_section:
        document.add_section()
        document.add_paragraph("second section text")
    out = BytesIO()
    document.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------


def pptx_bytes(
    slides: Sequence[tuple[str | None, Sequence[str]]],
    *,
    notes: str | None = None,
    table: Sequence[Sequence[str]] | None = None,
    picture: bool = False,
    grouped: str | None = None,
    nested: str | None = None,
    blank_notes: bool = False,
) -> bytes:
    """One slide per ``(title, body lines)``; the extras land on the first slide.

    *grouped* puts a text box inside a group; *nested* puts one inside a
    group inside a group; *blank_notes* attaches an empty notes slide.
    """
    presentation = pptx.Presentation()
    for index, (title, lines) in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1 if title is not None else 6])
        if title is not None:
            slide.shapes.title.text = title
            slide.placeholders[1].text_frame.text = "\n".join(lines)
        else:
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text_frame.text = "\n".join(lines)
        if index:
            continue
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
        if table is not None:
            grid = slide.shapes.add_table(len(table), len(table[0]), Inches(1), Inches(3), Inches(4), Inches(1)).table
            for r, row in enumerate(table):
                for c, cell in enumerate(row):
                    grid.cell(r, c).text = cell
        if picture:
            slide.shapes.add_picture(BytesIO(png_bytes()), Inches(1), Inches(5)).name = "Logo"
        if grouped is not None:
            group = slide.shapes.add_group_shape()
            box = group.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1))
            box.text_frame.text = grouped
        if nested is not None:
            inner = slide.shapes.add_group_shape().shapes.add_group_shape()
            box = inner.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1))
            box.text_frame.text = nested
        if blank_notes:
            slide.notes_slide.notes_text_frame.text = ""
    out = BytesIO()
    presentation.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


def xlsx_bytes(sheets: Mapping[str, Mapping[str, object]], *, hidden: Sequence[str] = ()) -> bytes:
    """One sheet per name holding the given ``cell address → value`` entries; *hidden* names are hidden."""
    workbook = openpyxl.Workbook()
    first = True
    for name, cells in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = name
        for address, value in cells.items():
            sheet[address] = value
        if name in hidden:
            sheet.sheet_state = "hidden"
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()
